"""Small helper layer over python-docx for the two GxP-style documents (Functional Spec, Test Protocol).

Styles: Calibri 11 body; Heading 1/2/3 in the accent colour; tables with a shaded header row that repeats on every page;
title page without header/footer; header with the document title; footer "Page X of Y"; a TOC field (updated by Word).
"""
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ACCENT_HEX = '1F4E79'
GREY = RGBColor(0x59, 0x59, 0x59)
FONT = 'Calibri'


def _set_font(style, size=None, bold=None, color=None):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), FONT)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def _field(paragraph, instr, placeholder=''):
    """Complex field (PAGE, NUMPAGES, TOC, SEQ). Word computes the value."""
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); run._r.append(b)
    run = paragraph.add_run()
    t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = ' ' + instr + ' '; run._r.append(t)
    run = paragraph.add_run()
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); run._r.append(s)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); run._r.append(e)


class Doc:
    def __init__(self, title, short_title):
        self.d = Document()
        self.title = title
        self.fig = 0
        self.tab = 0
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)       # US Letter
        sec.orientation = WD_ORIENT.PORTRAIT
        for side in ('left_margin', 'right_margin'):
            setattr(sec, side, Cm(2.2))
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
        sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)
        sec.different_first_page_header_footer = True
        st = self.d.styles
        _set_font(st['Normal'], 11)
        st['Normal'].paragraph_format.space_after = Pt(6)
        st['Normal'].paragraph_format.line_spacing = 1.08
        for name, size, before, after in (('Heading 1', 16, 18, 8), ('Heading 2', 13, 14, 6), ('Heading 3', 11.5, 10, 4)):
            _set_font(st[name], size, True, ACCENT)
            pf = st[name].paragraph_format
            pf.space_before, pf.space_after, pf.keep_with_next = Pt(before), Pt(after), True
        _set_font(st['Title'], 28, True, ACCENT)
        _set_font(st['Subtitle'], 14, False, GREY)
        _set_font(st['Caption'], 9.5, False, GREY)
        st['Caption'].font.italic = True
        _set_font(st['List Bullet'], 11)
        st['List Bullet'].paragraph_format.space_after = Pt(3)
        _set_font(st['List Number'], 11)
        st['List Number'].paragraph_format.space_after = Pt(3)
        for toc in ('TOC 1', 'TOC 2'):
            try:
                _set_font(st[toc], 11)
            except KeyError:
                pass
        # Header (document title) and footer (Page X of Y) for pages after the title page
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(short_title); r.font.size = Pt(9); r.font.color.rgb = GREY
        self._bottom_border(hp)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for part in ('Page ', ('PAGE', '1'), ' of ', ('NUMPAGES', '1')):
            if isinstance(part, tuple):
                _field(fp, part[0], part[1])
            else:
                fp.add_run(part)
        for r in fp.runs:
            r.font.size = Pt(9); r.font.color.rgb = GREY

    @staticmethod
    def _bottom_border(p):
        ppr = p._p.get_or_add_pPr()
        bdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:bottom')
        for k, v in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '4'), ('w:color', 'A6A6A6')):
            b.set(qn(k), v)
        bdr.append(b); ppr.append(bdr)

    # ---------------------------------------------------------------- blocks
    def title_page(self, title, subtitle, control_rows, revisions, approvals=None, top_blank=5):
        for _ in range(top_blank):
            self.d.add_paragraph()
        self.d.add_paragraph(title, style='Title')
        self.d.add_paragraph(subtitle, style='Subtitle')
        self.d.add_paragraph()
        self.para('Document control', bold=True, color=ACCENT, size=12, keep_next=True)
        self.table(['Field', 'Value'], control_rows, widths=[4.5, 12.6], caption=False)
        self.para('Revision history', bold=True, color=ACCENT, size=12, keep_next=True)
        self.table(['Version', 'Date', 'Author', 'Change'], revisions, widths=[2.0, 2.6, 3.8, 8.7], caption=False)
        if approvals:
            self.para('Approval', bold=True, color=ACCENT, size=12, keep_next=True)
            self.table(['Role', 'Name', 'Signature', 'Date'], approvals, widths=[4.0, 4.5, 5.0, 3.6], caption=False,
                       row_height=1.0)
        self.page_break()

    def toc(self):
        self.para('Contents', bold=True, color=ACCENT, size=16)
        p = self.d.add_paragraph()
        _field(p, 'TOC \\o "1-2" \\h \\z \\u', 'Right-click and choose Update Field to build the table of contents.')
        self.page_break()

    def h1(self, text, new_page=True):
        h = self.d.add_heading(text, level=1)
        h.paragraph_format.page_break_before = new_page
        return h

    def h2(self, text):
        return self.d.add_heading(text, level=2)

    def h3(self, text):
        return self.d.add_heading(text, level=3)

    def para(self, text='', bold=False, color=None, size=None, italic=False, keep_next=False, align=None):
        p = self.d.add_paragraph()
        self._runs(p, text, bold=bold, color=color, size=size, italic=italic)
        if keep_next:
            p.paragraph_format.keep_with_next = True
        if align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    @staticmethod
    def _runs(p, text, bold=False, color=None, size=None, italic=False):
        """Text with **bold** spans."""
        parts = text.split('**')
        for i, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run(part)
            r.bold = bold or (i % 2 == 1)
            r.italic = italic
            if color is not None:
                r.font.color.rgb = color
            if size:
                r.font.size = Pt(size)

    def bullets(self, items, style='List Bullet'):
        for it in items:
            p = self.d.add_paragraph(style=style)
            self._runs(p, it)

    def numbered(self, items):
        for n, it in enumerate(items, 1):
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent, pf.space_after = Cm(0.95), Cm(-0.6), Pt(3)
            self._runs(p, f'{n}.	' + it)
            pf.tab_stops.add_tab_stop(Cm(0.95))

    def landscape_begin(self):
        old = self.d.sections[-1]
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = old.page_height, old.page_width
        sec.different_first_page_header_footer = False

    def landscape_end(self):
        old = self.d.sections[-1]
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = old.page_height, old.page_width
        sec.different_first_page_header_footer = False

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def figure(self, path, width_cm, caption):
        self.fig += 1
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width_cm))
        c = self.d.add_paragraph(style='Caption')
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.add_run(f'Figure {self.fig}. {caption}')

    def table(self, header, rows, widths, caption=None, size=9.5, row_height=None):
        """Header row shaded in the accent colour, bold white, repeated on every page; rows don't split across pages."""
        if caption:
            self.tab += 1
            c = self.d.add_paragraph(style='Caption')
            c.add_run(f'Table {self.tab}. {caption}')
            c.paragraph_format.keep_with_next = True
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = self.d.styles['Table Grid']
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._table_borders(t)
        hdr = t.rows[0]
        trpr = hdr._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trpr.append(th)
        for i, h in enumerate(header):
            cell = hdr.cells[i]
            self._shade(cell, ACCENT_HEX)
            p = cell.paragraphs[0]
            r = p.add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            tr = t.rows[-1]._tr.get_or_add_trPr()
            cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); tr.append(cs)
            if row_height:
                from docx.shared import Cm as _Cm
                t.rows[-1].height = _Cm(row_height)
            for i, val in enumerate(row):
                lines = str(val).split('\n')
                p = cells[i].paragraphs[0]
                for li, line in enumerate(lines):
                    if li:
                        p = cells[i].add_paragraph()
                    self._runs(p, line, size=size)
                    p.paragraph_format.space_after = Pt(1)
                if ri % 2 == 1:
                    self._shade(cells[i], 'F2F6FA')
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        n = len(t.rows)
        keep_rows = range(0, n - 1) if n <= 12 else range(n - 3, n - 1)
        for ri in keep_rows:
            for cell in t.rows[ri].cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
        # small spacer after the table
        sp = self.d.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
        return t

    @staticmethod
    def _shade(cell, hex_fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
        tcpr.append(shd)

    @staticmethod
    def _table_borders(t):
        tblpr = t._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement(f'w:{edge}')
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'BFBFBF')
            borders.append(e)
        tblpr.append(borders)
        mar = OxmlElement('w:tblCellMar')
        for side, v in (('top', 40), ('bottom', 40), ('left', 80), ('right', 80)):
            e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
        tblpr.append(mar)

    def save(self, path):
        self.d.save(path)
